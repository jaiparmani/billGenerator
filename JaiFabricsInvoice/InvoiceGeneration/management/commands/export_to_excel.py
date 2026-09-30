import os

from django.apps import apps
from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

DEFAULT_OUTPUT_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__)))),
    "database_export.xlsx",
)


class Command(BaseCommand):
    help = "Export every model in the InvoiceGeneration app to a sheet in an Excel workbook."

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            default=DEFAULT_OUTPUT_PATH,
            help="Path to write the .xlsx file to (default: %(default)s)",
        )

    def handle(self, *args, **options):
        output_path = options["output"]
        workbook = Workbook()
        workbook.remove(workbook.active)

        models = apps.get_app_config("InvoiceGeneration").get_models()
        for model in models:
            sheet = workbook.create_sheet(title=model.__name__[:31])
            field_names = [field.name for field in model._meta.fields]
            sheet.append(field_names)

            for row in model.objects.all().values_list(*field_names):
                sheet.append([str(value) if value is not None else "" for value in row])

            for i, field_name in enumerate(field_names, start=1):
                width = max(len(field_name), 12)
                sheet.column_dimensions[get_column_letter(i)].width = min(width + 4, 40)

        workbook.save(output_path)
        self.stdout.write(self.style.SUCCESS(f"Exported database to {output_path}"))
